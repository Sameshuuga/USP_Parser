import shutil, logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import (
    coordinate_from_string,
    column_index_from_string,
    get_column_letter,
)

from main.settings import default_output_dir as dod

logger = logging.getLogger(__name__)


def copy_templet(
    templet_location: Path, copy_location: Path = dod / "temp.xlsx"
) -> Path:
    output = shutil.copy(templet_location, copy_location)
    logger.info(f"{templet_location.name} copied to {copy_location}")
    return output


def generate_cell_refs(payload: list, start_cell: str = "A1") -> list:
    """
    Auto-generate (cell_ref, value) pairs starting from start_cell,
    filling left-to-right, then top-to-bottom.

    payload: flat list of values (fills a single row),
             or list of lists (fills a 2D grid, one sub-list per row)
    start_cell: top-left cell to begin from, e.g. "A1", "C5"
    """
    col_letter, start_row = coordinate_from_string(start_cell)
    start_col = column_index_from_string(col_letter)

    # Normalize a flat list into a single row
    if not payload or not isinstance(payload[0], (list, tuple)):
        payload = [payload]

    refs = []
    for row_offset, row in enumerate(payload):
        for col_offset, value in enumerate(row):
            col = get_column_letter(start_col + col_offset)
            row_num = start_row + row_offset
            refs.append((f"{col}{row_num}", value))

    logger.info("Cell ref generation complete.")
    return refs


def inject_xlsx(
    target_file: Path, payload: list, start_cell: str = "A1", sheet_name: str = None
) -> Path:
    """
    Write values into an existing xlsx file starting from start_cell.

    target_file: path to the .xlsx file
    payload: flat list of values or list of lists (2D grid)
    start_cell: top-left cell to begin writing from (default "A1")
    sheet_name: which sheet to write to (defaults to the active sheet)
    """
    logger.info(f"Injecting {target_file} starting at {start_cell}")
    wb = load_workbook(target_file)
    ws = wb[sheet_name] if sheet_name else wb.active

    for cell_ref, value in generate_cell_refs(payload, start_cell):
        ws[cell_ref] = value

    wb.save(target_file)
    logger.info("Injection complete.")
    return target_file


if __name__ == "__main__":
    target_file = copy_templet(dod / "excel_templet.xlsx")
    payload = [["H", "e", "l", "l", "o"], ["N", "e", "w"], ["W", "o", "r", "l", "d"]]
    inject_xlsx(target_file, payload)
